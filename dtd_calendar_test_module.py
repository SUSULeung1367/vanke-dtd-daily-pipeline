"""Calendar-gated daily DTD test module.

The module reads confirmed history from ``vanke.xlsx``, reads incremental rows
from ``data_temporary.xlsx``, and appends successful results to a separate
``temporary_output.xlsx`` table with columns Comp_no, Date, and DTD.

Only dates marked HK_Open in the supplied calendar may produce DTD.  An open
date must also be the next Hong Kong trading session after the last completed
confirmed/temporary output date.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional, Union
import os
import uuid

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from scipy.optimize import brentq, minimize_scalar
from scipy.stats import norm


COLS = {
    "company": "Comp_no",
    "date": "Date",
    "equity": "CUR_MKT_CAP(HKD)",
    "cl": "BS_CUR_LIAB(HKD)",
    "ltd": "BS_LT_BORROW(HKD)",
    "tl": "BS_TOT_LIAB2(HKD)",
    "assets": "BS_TOT_ASSET(HKD)",
    "rf": "Risk_Free_Rate",
}
INPUT_COLUMNS = [
    COLS["company"],
    COLS["date"],
    COLS["equity"],
    COLS["cl"],
    COLS["ltd"],
    COLS["tl"],
    COLS["assets"],
    COLS["rf"],
]
MODEL_REQUIRED_COLUMNS = [
    COLS[key] for key in ("date", "equity", "cl", "ltd", "tl", "assets", "rf")
]
BS_COLUMNS = [COLS[key] for key in ("cl", "ltd", "tl", "assets")]
OUTPUT_COLUMNS = [COLS["company"], COLS["date"], "DTD"]
PathLike = Union[str, Path]


class DTDTestError(RuntimeError):
    """Base error for calendar-gated daily processing."""


class CalendarCoverageError(DTDTestError):
    """The requested or next required date is outside the supplied calendar."""


class MissingTradingDayError(DTDTestError):
    """A later trading day was requested before the required next session."""


class AlreadyProcessedError(DTDTestError):
    """The target trading day is not later than the completed output history."""


class InputDataError(DTDTestError):
    """The input tables are missing a required field or target-date row."""


class OutputIntegrityError(DTDTestError):
    """The temporary output history is duplicated, off-calendar, or discontinuous."""


@dataclass(frozen=True)
class DTDResult:
    date: pd.Timestamp
    dtd: float
    sigma: float
    asset_value: float
    default_point: float
    n_obs: int
    delta: float
    source_status: str


@dataclass(frozen=True)
class DailyProcessingResult:
    requested_date: pd.Timestamp
    status: str
    hk_open: bool
    last_completed_date: pd.Timestamp
    expected_next_open_date: pd.Timestamp
    dtd: Optional[float]
    wrote_output: bool
    message: str
    diagnostics: Optional[DTDResult] = None

    def as_dict(self) -> dict:
        item = asdict(self)
        for field in (
            "requested_date",
            "last_completed_date",
            "expected_next_open_date",
        ):
            value = item[field]
            item[field] = (
                None if pd.isna(value) else pd.Timestamp(value).strftime("%Y-%m-%d")
            )
        if self.diagnostics is not None:
            item["diagnostics"]["date"] = self.diagnostics.date.strftime("%Y-%m-%d")
        return item


def _parse_single_date(value) -> pd.Timestamp:
    text = str(value).strip().removesuffix(".0")
    if len(text) == 8 and text.isdigit():
        parsed = pd.to_datetime(text, format="%Y%m%d", errors="coerce")
    else:
        parsed = pd.to_datetime(text, errors="coerce")
    if pd.isna(parsed):
        raise InputDataError(f"Cannot parse date: {value!r}")
    return pd.Timestamp(parsed).normalize()


def _parse_dates(series: pd.Series) -> pd.Series:
    text = series.astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    compact = text.str.fullmatch(r"\d{8}")
    parsed = pd.to_datetime(text, errors="coerce")
    if compact.any():
        parsed.loc[compact] = pd.to_datetime(
            text.loc[compact], format="%Y%m%d", errors="coerce"
        )
    return parsed.dt.normalize()


def _as_boolean(series: pd.Series, label: str) -> pd.Series:
    if pd.api.types.is_bool_dtype(series):
        return series.fillna(False)
    mapping = {
        "true": True,
        "1": True,
        "yes": True,
        "y": True,
        "open": True,
        "false": False,
        "0": False,
        "no": False,
        "n": False,
        "closed": False,
    }
    converted = series.astype(str).str.strip().str.lower().map(mapping)
    if converted.isna().any():
        bad = series[converted.isna()].astype(str).unique().tolist()[:10]
        raise InputDataError(f"{label} contains invalid Boolean values: {bad}")
    return converted.astype(bool)


def _read_excel_sheet(
    path: PathLike, preferred_sheet: str, allow_missing_file: bool = False
) -> pd.DataFrame:
    path = Path(path)
    if not path.exists():
        if allow_missing_file:
            return pd.DataFrame()
        raise FileNotFoundError(f"File not found: {path}")
    if path.stat().st_size == 0:
        return pd.DataFrame()
    with pd.ExcelFile(path) as excel:
        if preferred_sheet in excel.sheet_names:
            selected = preferred_sheet
        elif len(excel.sheet_names) == 1:
            selected = excel.sheet_names[0]
        else:
            raise InputDataError(
                f"{path.name} has no sheet {preferred_sheet!r}; available sheets: {excel.sheet_names}"
            )
        return pd.read_excel(excel, sheet_name=selected)


def load_hk_calendar(
    calendar_file: PathLike, sheet_name: str = "Daily Calendar"
) -> pd.DataFrame:
    calendar = _read_excel_sheet(calendar_file, sheet_name)
    needed = {"Date", "HK_Open"}
    if not needed.issubset(calendar.columns):
        raise InputDataError(
            f"Trading calendar is missing columns: {sorted(needed - set(calendar.columns))}"
        )
    calendar = calendar.copy()
    calendar["Date"] = _parse_dates(calendar["Date"])
    if calendar["Date"].isna().any():
        raise InputDataError("Trading calendar contains an invalid Date")
    if calendar["Date"].duplicated().any():
        dates = (
            calendar.loc[calendar["Date"].duplicated(keep=False), "Date"]
            .dt.strftime("%Y-%m-%d")
            .tolist()
        )
        raise InputDataError(f"Trading calendar contains duplicate dates: {dates[:10]}")
    calendar["HK_Open"] = _as_boolean(calendar["HK_Open"], "HK_Open")
    return calendar.sort_values("Date").reset_index(drop=True)


def _validate_input_schema(frame: pd.DataFrame, label: str) -> None:
    missing = [
        column for column in MODEL_REQUIRED_COLUMNS if column not in frame.columns
    ]
    if missing:
        raise InputDataError(f"{label} is missing columns: {missing}")


def combine_input_history(
    confirmed_input: pd.DataFrame,
    temporary_input: pd.DataFrame,
    company: Optional[object] = None,
) -> pd.DataFrame:
    _validate_input_schema(confirmed_input, "vanke Input")
    _validate_input_schema(temporary_input, "data_temporary")
    confirmed = confirmed_input.copy()
    temporary = temporary_input.copy()
    confirmed["_source_status"] = "confirmed"
    temporary["_source_status"] = "temporary"
    data = pd.concat([confirmed, temporary], ignore_index=True, sort=False)
    data[COLS["date"]] = _parse_dates(data[COLS["date"]])
    if data[COLS["date"]].isna().any():
        raise InputDataError("Input contains an invalid Date")

    has_company = COLS["company"] in data.columns
    if company is not None:
        if not has_company:
            raise InputDataError(
                "company was provided, but Input has no Comp_no column"
            )
        data = data[data[COLS["company"]] == company].copy()
        if data.empty:
            raise InputDataError(f"No Input rows found for Comp_no={company}")
    elif has_company:
        companies = data[COLS["company"]].dropna().unique()
        if len(companies) > 1:
            raise InputDataError(
                f"Input contains multiple companies {companies.tolist()}; provide company"
            )

    priority = data["_source_status"].map({"temporary": 1, "confirmed": 2})
    data = data.assign(_source_priority=priority)
    keys = ([COLS["company"]] if has_company else []) + [COLS["date"]]
    data = (
        data.sort_values(keys + ["_source_priority"])
        .drop_duplicates(keys, keep="last")
        .sort_values(keys)
        .reset_index(drop=True)
    )
    for column in [COLS[key] for key in ("equity", "cl", "ltd", "tl", "assets", "rf")]:
        data[column] = pd.to_numeric(data[column], errors="coerce")
    if has_company:
        data[BS_COLUMNS] = data.groupby(COLS["company"], dropna=False)[
            BS_COLUMNS
        ].ffill()
    else:
        data[BS_COLUMNS] = data[BS_COLUMNS].ffill()
    if data[BS_COLUMNS].isna().any().any():
        raise InputDataError(
            "Financial fields cannot be carried forward; the first history row must be complete"
        )
    return data.drop(columns="_source_priority")


def _normalize_output(frame: pd.DataFrame, label: str) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=OUTPUT_COLUMNS)
    missing = [column for column in OUTPUT_COLUMNS if column not in frame.columns]
    if missing:
        raise OutputIntegrityError(f"{label} is missing columns: {missing}")
    output = frame[OUTPUT_COLUMNS].copy()
    output[COLS["date"]] = _parse_dates(output[COLS["date"]])
    if output[COLS["date"]].isna().any():
        raise OutputIntegrityError(f"{label} contains an invalid Date")
    return output


def _rate_to_decimal(value: float) -> float:
    value = float(value)
    return value / 100.0 if abs(value) > 1.0 else value


def add_model_fields(frame: pd.DataFrame, delta: float = 0.5) -> pd.DataFrame:
    if not 0.0 <= delta <= 1.0:
        raise InputDataError("delta must be between 0 and 1")
    data = frame.copy().sort_values(COLS["date"]).reset_index(drop=True)
    data["Other_Liabilities(HKD)"] = (
        data[COLS["tl"]] - data[COLS["cl"]] - data[COLS["ltd"]]
    )
    data["Default_Point(HKD)"] = (
        data[COLS["cl"]]
        + 0.5 * data[COLS["ltd"]]
        + delta * data["Other_Liabilities(HKD)"]
    )
    data["Risk_Free_Decimal"] = data[COLS["rf"]].map(_rate_to_decimal)
    return data


def implied_asset_value(
    equity: float,
    debt: float,
    rate: float,
    sigma: float,
    maturity: float = 1.0,
) -> float:
    equity, debt, rate, sigma, maturity = map(
        float, (equity, debt, rate, sigma, maturity)
    )
    if equity <= 0 or debt <= 0 or sigma <= 0 or maturity <= 0:
        raise InputDataError(
            "equity, default point, sigma and maturity must be positive"
        )
    sqrt_t = np.sqrt(maturity)

    def residual(asset_value: float) -> float:
        d1 = (np.log(asset_value / debt) + (rate + 0.5 * sigma**2) * maturity) / (
            sigma * sqrt_t
        )
        d2 = d1 - sigma * sqrt_t
        return (
            asset_value * norm.cdf(d1)
            - debt * np.exp(-rate * maturity) * norm.cdf(d2)
            - equity
        )

    lower = max(equity, 1e-9)
    upper = equity + debt * np.exp(-rate * maturity) + debt
    while residual(upper) < 0:
        upper *= 2.0
        if upper > 1e12 * max(1.0, debt):
            raise DTDTestError(
                "Cannot build a root-solving interval for market asset value"
            )
    return float(brentq(residual, lower, upper, maxiter=200))


def estimate_sigma(
    history: pd.DataFrame,
    as_of: pd.Timestamp,
    delta: float = 0.5,
    lookback_days: int = 365,
    min_obs: int = 50,
    trading_days_per_year: int = 250,
    sigma_bounds: tuple[float, float] = (0.005, 1.50),
) -> tuple[float, pd.DataFrame]:
    data = add_model_fields(history, delta)
    start = as_of - pd.Timedelta(days=lookback_days)
    window = data[(data[COLS["date"]] > start) & (data[COLS["date"]] <= as_of)].copy()
    needed = MODEL_REQUIRED_COLUMNS + ["Default_Point(HKD)", "Risk_Free_Decimal"]
    window = window.dropna(subset=needed)
    window = window[
        (window[COLS["equity"]] > 0)
        & (window[COLS["assets"]] > 0)
        & (window["Default_Point(HKD)"] > 0)
    ].reset_index(drop=True)
    if len(window) < min_obs:
        raise InputDataError(
            f"At least {min_obs} valid observations are required; found {len(window)}"
        )

    equity = window[COLS["equity"]].to_numpy(float)
    debt = window["Default_Point(HKD)"].to_numpy(float)
    book_assets = window[COLS["assets"]].to_numpy(float)
    rates = window["Risk_Free_Decimal"].to_numpy(float)
    h = np.full(len(window) - 1, 1.0 / trading_days_per_year)

    def negative_log_likelihood(sigma: float) -> float:
        try:
            market_assets = np.array(
                [
                    implied_asset_value(equity[i], debt[i], rates[i], sigma)
                    for i in range(len(window))
                ]
            )
            d1 = (np.log(market_assets / debt) + rates + 0.5 * sigma**2) / sigma
            normal_d1 = np.clip(norm.cdf(d1), 1e-14, 1.0)
            transformed_return = np.log(
                (market_assets[1:] / book_assets[1:])
                * (book_assets[:-1] / market_assets[:-1])
            )
            mu = 0.5 * sigma**2 + transformed_return.sum() / h.sum()
            residuals = transformed_return - (mu - 0.5 * sigma**2) * h
            log_likelihood = (
                -0.5 * (len(window) - 1) * np.log(2 * np.pi)
                - 0.5 * np.sum(np.log(sigma**2 * h))
                - np.sum(np.log(market_assets[1:] / book_assets[1:]))
                - np.sum(np.log(normal_d1[1:]))
                - 0.5 / sigma**2 * np.sum((residuals**2) / h)
            )
            return -float(log_likelihood) if np.isfinite(log_likelihood) else 1e100
        except (ValueError, RuntimeError, FloatingPointError):
            return 1e100

    optimum = minimize_scalar(
        negative_log_likelihood,
        bounds=sigma_bounds,
        method="bounded",
        options={"xatol": 1e-8},
    )
    if not optimum.success or not np.isfinite(optimum.fun):
        raise DTDTestError("Sigma optimization did not converge")
    return float(optimum.x), window


def calculate_dtd(
    history: pd.DataFrame,
    as_of: pd.Timestamp,
    delta: float = 0.5,
    lookback_days: int = 365,
    min_obs: int = 50,
    trading_days_per_year: int = 250,
) -> DTDResult:
    data = add_model_fields(history, delta)
    data = data[data[COLS["date"]] <= as_of].copy()
    target_rows = data[data[COLS["date"]] == as_of]
    if target_rows.empty:
        raise InputDataError(
            f"Temporary Input has no row for target trading date {as_of.strftime('%Y%m%d')}"
        )
    row = target_rows.iloc[-1]
    sigma, window = estimate_sigma(
        data,
        as_of,
        delta=delta,
        lookback_days=lookback_days,
        min_obs=min_obs,
        trading_days_per_year=trading_days_per_year,
    )
    market_assets = implied_asset_value(
        row[COLS["equity"]], row["Default_Point(HKD)"], row["Risk_Free_Decimal"], sigma
    )
    dtd = np.log(market_assets / row["Default_Point(HKD)"]) / sigma
    return DTDResult(
        date=as_of,
        dtd=float(dtd),
        sigma=float(sigma),
        asset_value=float(market_assets),
        default_point=float(row["Default_Point(HKD)"]),
        n_obs=len(window),
        delta=float(delta),
        source_status=str(row.get("_source_status", "unknown")),
    )


def _calendar_row(calendar: pd.DataFrame, date: pd.Timestamp) -> pd.Series:
    row = calendar[calendar["Date"] == date]
    if row.empty:
        start = calendar["Date"].min().strftime("%Y-%m-%d")
        end = calendar["Date"].max().strftime("%Y-%m-%d")
        raise CalendarCoverageError(
            f"Date {date.date()} is outside calendar coverage {start} to {end}"
        )
    return row.iloc[0]


def _next_hk_open(calendar: pd.DataFrame, after_date: pd.Timestamp) -> pd.Timestamp:
    candidates = calendar[(calendar["Date"] > after_date) & calendar["HK_Open"]]
    if candidates.empty:
        end = calendar["Date"].max().strftime("%Y-%m-%d")
        raise CalendarCoverageError(
            f"Calendar has no open date after {after_date.date()}; it ends on {end}"
        )
    return pd.Timestamp(candidates.iloc[0]["Date"])


def _select_company(
    frame: pd.DataFrame, company: Optional[object], label: str
) -> pd.DataFrame:
    if frame.empty or COLS["company"] not in frame.columns:
        return frame.copy()
    if company is not None:
        return frame[frame[COLS["company"]] == company].copy()
    companies = frame[COLS["company"]].dropna().unique()
    if len(companies) > 1:
        raise InputDataError(
            f"{label} contains multiple companies {companies.tolist()}; provide company"
        )
    return frame.copy()


def _validate_output_history(
    confirmed_output: pd.DataFrame,
    temporary_output: pd.DataFrame,
    calendar: pd.DataFrame,
    company: Optional[object],
) -> tuple[pd.Timestamp, object]:
    confirmed = _select_company(
        _normalize_output(confirmed_output, "vanke Output"), company, "vanke Output"
    )
    temporary = _select_company(
        _normalize_output(temporary_output, "temporary_output"),
        company,
        "temporary_output",
    )
    if confirmed.empty:
        raise OutputIntegrityError(
            "Confirmed Vanke Output is empty; production anchor is unavailable"
        )
    company_value = company
    if company_value is None and COLS["company"] in confirmed.columns:
        values = confirmed[COLS["company"]].dropna().unique()
        company_value = values[0] if len(values) == 1 else None

    confirmed_last = confirmed[COLS["date"]].max()
    if temporary[COLS["date"]].duplicated().any():
        duplicates = temporary.loc[
            temporary[COLS["date"]].duplicated(keep=False), COLS["date"]
        ]
        raise OutputIntegrityError(
            "Temporary Output contains duplicate dates: "
            + ", ".join(duplicates.dt.strftime("%Y-%m-%d").unique())
        )
    if not temporary.empty:
        temp_calendar = temporary[[COLS["date"]]].merge(
            calendar[["Date", "HK_Open"]],
            left_on=COLS["date"],
            right_on="Date",
            how="left",
        )
        invalid = temp_calendar[temp_calendar["HK_Open"] != True]  # noqa: E712
        if not invalid.empty:
            dates = invalid[COLS["date"]].dt.strftime("%Y-%m-%d").tolist()
            raise OutputIntegrityError(
                f"Temporary Output contains closed or off-calendar dates: {dates}"
            )
        after_anchor = temporary[temporary[COLS["date"]] > confirmed_last]
        if not after_anchor.empty:
            temporary_last = after_anchor[COLS["date"]].max()
            expected = set(
                calendar.loc[
                    (calendar["Date"] > confirmed_last)
                    & (calendar["Date"] <= temporary_last)
                    & calendar["HK_Open"],
                    "Date",
                ]
            )
            actual = set(after_anchor[COLS["date"]])
            missing = sorted(expected - actual)
            if missing:
                missing_text = ", ".join(date.strftime("%Y-%m-%d") for date in missing)
                raise OutputIntegrityError(
                    f"Temporary Output is not continuous; missing: {missing_text}"
                )
    completed = (
        confirmed
        if temporary.empty
        else pd.concat([confirmed, temporary], ignore_index=True, sort=False)
    )
    return pd.Timestamp(completed[COLS["date"]].max()), company_value


def _append_temporary_output(
    output_file: PathLike,
    company_value: object,
    date: pd.Timestamp,
    dtd: float,
    sheet_name: str = "Output",
) -> None:
    path = Path(output_file)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and path.stat().st_size > 0:
        workbook = load_workbook(path)
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        elif len(workbook.sheetnames) == 1 and workbook.active.max_row <= 1:
            sheet = workbook.active
            sheet.title = sheet_name
        else:
            sheet = workbook.create_sheet(sheet_name)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name

    existing_headers = [sheet.cell(1, column).value for column in range(1, 4)]
    if sheet.max_row == 1 and all(value is None for value in existing_headers):
        for column, header in enumerate(OUTPUT_COLUMNS, start=1):
            sheet.cell(1, column, header)
    elif existing_headers != OUTPUT_COLUMNS:
        raise OutputIntegrityError(
            f"{path.name} sheet {sheet_name} must use headers {OUTPUT_COLUMNS}; found {existing_headers}"
        )

    date_int = int(date.strftime("%Y%m%d"))
    for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        if row[1] is not None and _parse_single_date(row[1]) == date:
            raise AlreadyProcessedError(
                f"Temporary Output already contains {date.date()}; it will not be overwritten"
            )
    sheet.append([company_value, date_int, float(dtd)])

    temp_path = path.with_name(
        f".{path.stem}.{uuid.uuid4().hex}.tmp{path.suffix or '.xlsx'}"
    )
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def process_daily_dtd(
    target_date,
    *,
    confirmed_file: PathLike,
    temporary_input_file: PathLike,
    temporary_output_file: PathLike,
    calendar_file: PathLike,
    company: Optional[object] = None,
    confirmed_input_sheet: str = "Input",
    confirmed_output_sheet: str = "Output",
    temporary_input_sheet: str = "Input",
    temporary_output_sheet: str = "Output",
    calendar_sheet: str = "Daily Calendar",
    delta: float = 0.5,
    lookback_days: int = 365,
    min_obs: int = 50,
    trading_days_per_year: int = 250,
    write_output: bool = True,
) -> DailyProcessingResult:
    """Validate one date, calculate DTD when eligible, and append temporary output.

    Closed Hong Kong dates return ``SKIPPED_HK_CLOSED`` without writing.  An
    eligible date must be the next HK-open session after the latest completed
    confirmed/temporary output date.  Later dates raise MissingTradingDayError.
    """
    target = _parse_single_date(target_date)
    calendar = load_hk_calendar(calendar_file, calendar_sheet)
    target_calendar = _calendar_row(calendar, target)

    confirmed_output = _read_excel_sheet(confirmed_file, confirmed_output_sheet)
    temporary_output = _read_excel_sheet(
        temporary_output_file, temporary_output_sheet, allow_missing_file=True
    )
    last_completed, company_value = _validate_output_history(
        confirmed_output, temporary_output, calendar, company
    )
    expected = _next_hk_open(calendar, last_completed)

    if not bool(target_calendar["HK_Open"]):
        reason = target_calendar.get("Closed_Reason", "HK market closed")
        if pd.isna(reason) or not str(reason).strip():
            reason = "HK market closed"
        return DailyProcessingResult(
            requested_date=target,
            status="SKIPPED_HK_CLOSED",
            hk_open=False,
            last_completed_date=last_completed,
            expected_next_open_date=expected,
            dtd=None,
            wrote_output=False,
            message=f"HK is closed on {target.date()} ({reason}); no DTD is calculated or written.",
        )

    if target <= last_completed:
        raise AlreadyProcessedError(
            f"Target date {target.date()} is not later than completed date {last_completed.date()}; history will not be overwritten"
        )
    if target != expected:
        raise MissingTradingDayError(
            f"Cannot process {target.date()}; last completed date is {last_completed.date()}. "
            f"Process the next HK trading date {expected.date()} first"
        )

    confirmed_input = _read_excel_sheet(confirmed_file, confirmed_input_sheet)
    temporary_input = _read_excel_sheet(temporary_input_file, temporary_input_sheet)
    history = combine_input_history(confirmed_input, temporary_input, company)
    if not (history[COLS["date"]] == target).any():
        raise InputDataError(
            f"Temporary Input has no row for {target.strftime('%Y%m%d')}"
        )
    diagnostics = calculate_dtd(
        history,
        target,
        delta=delta,
        lookback_days=lookback_days,
        min_obs=min_obs,
        trading_days_per_year=trading_days_per_year,
    )
    if diagnostics.source_status != "temporary":
        raise InputDataError(
            f"Target date {target.date()} is not a Temporary Input row; confirmed dates are not recalculated"
        )

    if write_output:
        _append_temporary_output(
            temporary_output_file,
            company_value,
            target,
            diagnostics.dtd,
            temporary_output_sheet,
        )
    return DailyProcessingResult(
        requested_date=target,
        status="DTD_WRITTEN" if write_output else "DTD_CALCULATED_NOT_WRITTEN",
        hk_open=True,
        last_completed_date=last_completed,
        expected_next_open_date=expected,
        dtd=diagnostics.dtd,
        wrote_output=write_output,
        message=(
            f"{target.date()} is the next HK trading date; DTD={diagnostics.dtd:.10f}. "
            + (
                "Written to Temporary Output."
                if write_output
                else "No file was written."
            )
        ),
        diagnostics=diagnostics,
    )


def read_temporary_output(
    output_file: PathLike, sheet_name: str = "Output"
) -> pd.DataFrame:
    """Read the Date/DTD result table for display or downstream checks."""
    output = _read_excel_sheet(output_file, sheet_name, allow_missing_file=True)
    if output.empty:
        return pd.DataFrame(columns=OUTPUT_COLUMNS)
    normalized = _normalize_output(output, "temporary_output")
    normalized[COLS["date"]] = (
        normalized[COLS["date"]].dt.strftime("%Y%m%d").astype(int)
    )
    return normalized
