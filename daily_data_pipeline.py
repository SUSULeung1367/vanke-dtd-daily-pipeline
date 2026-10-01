"""Daily market-data ingestion and Input quality control for the Vanke demo.

The public entry point is ``process_daily_data``.  It handles one calendar
date, records source/QC evidence and writes only an eligible Temporary Input
row.  It never writes to confirmed Input or confirmed DTD Output.
"""

from pathlib import Path
from copy import copy
from datetime import datetime
import os
import time
import uuid
import warnings

import pandas as pd
import requests
import yfinance as yf
from openpyxl import load_workbook


# ============================================================
# MODULE 1 - CONFIGURATION
# ============================================================

PROJECT_DIR = Path(
    os.environ.get("DTD_PROJECT_DIR", Path(__file__).resolve().parent)
).resolve()

CORE_FILE = PROJECT_DIR / "vanke.xlsx"
CALENDAR_FILE = PROJECT_DIR / "China_HK_Trading_Calendar.xlsx"
ISSUED_CAPITAL_FILE = PROJECT_DIR / "Vanke Issued Capital DataLog.xlsx"
RISK_FREE_FILE = PROJECT_DIR / "HKMA_Risk_Free_Daily.xlsx"

DAILY_DATALOG_FILE = PROJECT_DIR / "Vanke_Daily_Datalog.xlsx"
TEMPORARY_INPUT_FILE = PROJECT_DIR / "vanke_dtd_temporary_data.xlsx"

CHINA_TICKER = "000002.SZ"
HK_TICKER = "2202.HK"
FX_TICKER = "CNYHKD=X"

HKMA_URL = (
    "https://api.hkma.gov.hk/public/market-data-and-statistics/"
    "monthly-statistical-bulletin/efbn/efbn-yield-daily"
)

MAX_ATTEMPTS = 5
RETRY_DELAYS_SECONDS = [0, 1, 2, 4, 8]
REQUEST_TIMEOUT = (4, 8)

# Plausibility and manual-review thresholds.
FX_MIN = 0.50
FX_MAX = 2.00
RISK_FREE_MIN = -5.00
RISK_FREE_MAX = 20.00
MAX_PRICE_CHANGE = 0.30
MAX_FX_CHANGE = 0.05
MAX_RISK_FREE_CHANGE_PP = 1.00
MAX_MARKET_CAP_CHANGE = 0.30

TEMPORARY_COLUMNS = [
    "Comp_no",
    "Date",
    "CUR_MKT_CAP(HKD)",
    "BS_CUR_LIAB(HKD)",
    "BS_LT_BORROW(HKD)",
    "BS_TOT_LIAB2(HKD)",
    "BS_TOT_ASSET(HKD)",
    "Risk_Free_Rate",
]

BALANCE_SHEET_COLUMNS = [
    "BS_CUR_LIAB(HKD)",
    "BS_LT_BORROW(HKD)",
    "BS_TOT_LIAB2(HKD)",
    "BS_TOT_ASSET(HKD)",
]

YFINANCE_CACHE_DIR = PROJECT_DIR / ".yfinance_cache"
YFINANCE_CACHE_DIR.mkdir(parents=True, exist_ok=True)
yf.set_tz_cache_location(str(YFINANCE_CACHE_DIR))


def configure_project(project_dir):
    """Point all file-based stages at one project directory.

    This keeps the calculation code unchanged while allowing the teacher test
    to run against an isolated copy instead of the real working files.
    """
    global PROJECT_DIR, CORE_FILE, CALENDAR_FILE, ISSUED_CAPITAL_FILE
    global RISK_FREE_FILE, DAILY_DATALOG_FILE, TEMPORARY_INPUT_FILE
    global YFINANCE_CACHE_DIR

    PROJECT_DIR = Path(project_dir).expanduser().resolve()
    CORE_FILE = PROJECT_DIR / "vanke.xlsx"
    CALENDAR_FILE = PROJECT_DIR / "China_HK_Trading_Calendar.xlsx"
    ISSUED_CAPITAL_FILE = PROJECT_DIR / "Vanke Issued Capital DataLog.xlsx"
    RISK_FREE_FILE = PROJECT_DIR / "HKMA_Risk_Free_Daily.xlsx"
    DAILY_DATALOG_FILE = PROJECT_DIR / "Vanke_Daily_Datalog.xlsx"
    TEMPORARY_INPUT_FILE = PROJECT_DIR / "vanke_dtd_temporary_data.xlsx"
    YFINANCE_CACHE_DIR = PROJECT_DIR / ".yfinance_cache"
    YFINANCE_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    yf.set_tz_cache_location(str(YFINANCE_CACHE_DIR))
    return PROJECT_DIR


# ============================================================
# MODULE 2 - API RETRY AND ONLINE INGESTION
# ============================================================


def run_with_retry(source_name, data_date, fetch_once):
    """Run one API function up to MAX_ATTEMPTS and audit every attempt."""
    attempts = []
    last_error = None

    for attempt_number in range(1, MAX_ATTEMPTS + 1):
        delay = RETRY_DELAYS_SECONDS[attempt_number - 1]
        if delay:
            time.sleep(delay)

        started_at = datetime.now()

        try:
            value = fetch_once()
            attempts.append(
                {
                    "API": source_name,
                    "Attempt": attempt_number,
                    "Attempt_Status": "SUCCESS",
                    "Attempt_Time": started_at,
                    "Error_Type": None,
                    "Error_Message": None,
                }
            )
            return value, attempts, None

        except Exception as error:
            last_error = error
            attempts.append(
                {
                    "API": source_name,
                    "Attempt": attempt_number,
                    "Attempt_Status": "FAILED",
                    "Attempt_Time": started_at,
                    "Error_Type": type(error).__name__,
                    "Error_Message": str(error)[:500],
                }
            )

    return None, attempts, last_error


def _extract_close_series(downloaded, ticker):
    close = downloaded["Close"]
    if isinstance(close, pd.DataFrame):
        close = close[ticker] if ticker in close.columns else close.iloc[:, 0]

    close.index = pd.to_datetime(close.index)
    if close.index.tz is not None:
        close.index = close.index.tz_localize(None)
    close.index = close.index.normalize()
    return pd.to_numeric(close, errors="coerce").dropna().sort_index()


def get_online_close(source_name, ticker, data_date, require_exact):
    """Download the latest Close and its preceding observation."""
    data_date = pd.Timestamp(data_date).normalize()

    def fetch_once():
        downloaded = yf.download(
            ticker,
            start=(data_date - pd.Timedelta(days=20)).strftime("%Y-%m-%d"),
            end=(data_date + pd.Timedelta(days=1)).strftime("%Y-%m-%d"),
            auto_adjust=False,
            progress=False,
            threads=False,
            timeout=REQUEST_TIMEOUT[1],
        )
        if downloaded.empty:
            raise ValueError(f"No online data returned for {ticker}.")

        close = _extract_close_series(downloaded, ticker)
        available = close.loc[close.index <= data_date]
        if available.empty:
            raise ValueError(f"No Close exists on or before {data_date.date()}.")

        source_date = pd.Timestamp(available.index[-1]).normalize()
        value = float(available.iloc[-1])
        previous = available.iloc[:-1]
        previous_value = float(previous.iloc[-1]) if not previous.empty else None
        previous_date = (
            pd.Timestamp(previous.index[-1]).normalize() if not previous.empty else None
        )
        return {
            "Value": value,
            "Source_Date": source_date,
            "Previous_Value": previous_value,
            "Previous_Source_Date": previous_date,
        }

    payload, attempts, error = run_with_retry(source_name, data_date, fetch_once)

    if payload is None:
        return {
            "OK": False,
            "Value": None,
            "Source_Date": None,
            "Previous_Value": None,
            "Previous_Source_Date": None,
            "Exact_Required": require_exact,
            "Status": f"FAILED_AFTER_{MAX_ATTEMPTS}_ATTEMPTS",
            "Error": str(error),
            "Attempts": attempts,
        }

    exact_match = payload["Source_Date"] == data_date
    return {
        "OK": (not require_exact) or exact_match,
        **payload,
        "Exact_Required": require_exact,
        "Status": "ONLINE_EXACT" if exact_match else "ONLINE_ASOF_PRIOR_DATE",
        "Error": (
            None if ((not require_exact) or exact_match) else "Source date mismatch"
        ),
        "Attempts": attempts,
    }


def load_risk_free_cache():
    if not RISK_FREE_FILE.exists():
        return pd.DataFrame(
            columns=["Date", "Day", "Risk_Free_Rate", "Risk_Free_Decimal"]
        )
    cache = pd.read_excel(RISK_FREE_FILE)
    if cache.empty:
        return cache
    cache["Date"] = pd.to_datetime(cache["Date"]).dt.normalize()
    cache["Risk_Free_Rate"] = pd.to_numeric(cache["Risk_Free_Rate"], errors="coerce")
    return cache.dropna(subset=["Date", "Risk_Free_Rate"])


def save_risk_free_cache(cache, data_date, value):
    new_row = pd.DataFrame(
        [
            {
                "Date": data_date,
                "Day": pd.Timestamp(data_date).day_name(),
                "Risk_Free_Rate": float(value),
                "Risk_Free_Decimal": float(value) / 100,
            }
        ]
    )
    updated = pd.concat([cache, new_row], ignore_index=True)
    updated["Date"] = pd.to_datetime(updated["Date"]).dt.normalize()
    updated = (
        updated.sort_values("Date")
        .drop_duplicates(subset=["Date"], keep="last")
        .reset_index(drop=True)
    )
    updated.to_excel(RISK_FREE_FILE, index=False)


def get_hkma_risk_free(data_date, require_exact, update_cache):
    """Use an exact cached/online rate for HK production days."""
    data_date = pd.Timestamp(data_date).normalize()
    cache = load_risk_free_cache()
    available = cache.loc[cache["Date"] <= data_date].sort_values("Date")
    exact = available.loc[available["Date"] == data_date]
    previous = available.loc[available["Date"] < data_date]
    previous_value = (
        float(previous.iloc[-1]["Risk_Free_Rate"]) if not previous.empty else None
    )

    if not exact.empty:
        row = exact.iloc[-1]
        attempts = [
            {
                "API": "HKMA_RISK_FREE",
                "Attempt": 0,
                "Attempt_Status": "CACHE_HIT",
                "Attempt_Time": datetime.now(),
                "Error_Type": None,
                "Error_Message": None,
            }
        ]
        return {
            "OK": True,
            "Value": float(row["Risk_Free_Rate"]),
            "Source_Date": pd.Timestamp(row["Date"]).normalize(),
            "Previous_Value": previous_value,
            "Status": "LOCAL_CACHE_EXACT",
            "Error": None,
            "Attempts": attempts,
        }

    if not require_exact:
        if available.empty:
            return {
                "OK": False,
                "Value": None,
                "Source_Date": None,
                "Previous_Value": None,
                "Status": "NO_CACHED_ASOF_VALUE",
                "Error": "No cached HKMA value",
                "Attempts": [],
            }
        row = available.iloc[-1]
        return {
            "OK": True,
            "Value": float(row["Risk_Free_Rate"]),
            "Source_Date": pd.Timestamp(row["Date"]).normalize(),
            "Previous_Value": previous_value,
            "Status": "LOCAL_CACHE_ASOF_CLOSED_DAY",
            "Error": None,
            "Attempts": [],
        }

    def fetch_once():
        response = requests.get(
            HKMA_URL,
            params={
                "choose": "end_of_day",
                "from": data_date.strftime("%Y-%m-%d"),
                "to": data_date.strftime("%Y-%m-%d"),
                "fields": "end_of_day,efb_364d",
                "pagesize": 5,
            },
            timeout=REQUEST_TIMEOUT,
        )
        response.raise_for_status()
        payload = response.json()
        if not payload.get("header", {}).get("success", False):
            raise ValueError(f"HKMA API error: {payload.get('header')}")
        online = pd.DataFrame(payload.get("result", {}).get("records", []))
        if online.empty:
            raise ValueError("HKMA returned no exact-date record.")
        online["Date"] = pd.to_datetime(online["end_of_day"])
        online["Value"] = pd.to_numeric(online["efb_364d"], errors="coerce")
        online = online.dropna(subset=["Date", "Value"])
        online = online.loc[online["Date"] == data_date]
        if online.empty:
            raise ValueError("HKMA returned no usable exact-date yield.")
        return float(online.iloc[-1]["Value"])

    value, attempts, error = run_with_retry("HKMA_RISK_FREE", data_date, fetch_once)

    if value is not None:
        if update_cache:
            save_risk_free_cache(cache, data_date, value)
        return {
            "OK": True,
            "Value": value,
            "Source_Date": data_date,
            "Previous_Value": previous_value,
            "Status": (
                "ONLINE_EXACT_AND_CACHED" if update_cache else "ONLINE_EXACT_DRY_RUN"
            ),
            "Error": None,
            "Attempts": attempts,
        }

    fallback_value = (
        float(available.iloc[-1]["Risk_Free_Rate"]) if not available.empty else None
    )
    fallback_date = (
        pd.Timestamp(available.iloc[-1]["Date"]).normalize()
        if not available.empty
        else None
    )
    return {
        "OK": False,
        "Value": fallback_value,
        "Source_Date": fallback_date,
        "Previous_Value": previous_value,
        "Status": f"FAILED_AFTER_{MAX_ATTEMPTS}_ATTEMPTS_STALE_CACHE_ONLY",
        "Error": str(error),
        "Attempts": attempts,
    }


# ============================================================
# MODULE 3 - DATA CHECKING AND QUALITY GATE
# ============================================================


def percentage_change(current, previous):
    if current is None or previous in (None, 0) or pd.isna(previous):
        return None
    return abs(float(current) / float(previous) - 1)


def check_daily_data(record, sources):
    bugs = []
    warnings_list = []
    review_items = []

    china = sources["CHINA_CLOSE"]
    hk = sources["HK_CLOSE"]
    fx = sources["FX"]
    risk_free = sources["RISK_FREE"]
    data_date = record["Date"]

    # ----- Blocking availability and source-date checks -----
    if not china["OK"]:
        bugs.append(f"China Close invalid: {china['Status']} {china['Error']}")
    if not hk["OK"]:
        bugs.append(f"HK Close invalid: {hk['Status']} {hk['Error']}")
    if not fx["OK"]:
        bugs.append(f"FX invalid: {fx['Status']} {fx['Error']}")

    if record["HK_Open"] and not risk_free["OK"]:
        bugs.append(f"Exact HKMA risk-free rate unavailable: {risk_free['Status']}")

    if record["China_Open"] and china["Source_Date"] != data_date:
        bugs.append("China market is open but China Close is not exact-date.")
    if record["HK_Open"] and hk["Source_Date"] != data_date:
        bugs.append("HK market is open but HK Close is not exact-date.")
    if record["HK_Open"] and fx["Source_Date"] != data_date:
        bugs.append("HK production date requires an exact-date FX observation.")
    if record["HK_Open"] and risk_free["Source_Date"] != data_date:
        bugs.append("HK production date requires an exact-date risk-free rate.")

    # ----- Closed-market carry-forward warnings -----
    if not record["China_Open"] and china["Value"] is not None:
        warnings_list.append(
            f"China closed; carried Close from {china['Source_Date'].date()}."
        )
    if not record["HK_Open"] and hk["Value"] is not None:
        warnings_list.append(
            f"HK closed; carried Close from {hk['Source_Date'].date()}."
        )
    if not record["HK_Open"] and risk_free["Source_Date"] is not None:
        warnings_list.append(
            f"HK closed; carried risk-free rate from {risk_free['Source_Date'].date()}."
        )

    # ----- Retry warnings -----
    for name, source in sources.items():
        failed_attempts = sum(
            row["Attempt_Status"] == "FAILED" for row in source["Attempts"]
        )
        if failed_attempts and source["OK"]:
            warnings_list.append(
                f"{name} succeeded after {failed_attempts} failed attempt(s)."
            )

    # ----- Hard numeric checks -----
    for label in ["China_Close_Used", "HK_Close_Used", "FX_Rate"]:
        value = record.get(label)
        if value is None or pd.isna(value) or float(value) <= 0:
            bugs.append(f"{label} must be present and positive.")

    if record["FX_Rate"] is not None and not FX_MIN <= record["FX_Rate"] <= FX_MAX:
        bugs.append(f"FX_Rate outside [{FX_MIN}, {FX_MAX}].")

    if record["Risk_Free_Rate"] is not None:
        if not RISK_FREE_MIN <= record["Risk_Free_Rate"] <= RISK_FREE_MAX:
            bugs.append("Risk_Free_Rate is outside the plausible range.")

    for label in ["China_Shares", "HK_Shares", *BALANCE_SHEET_COLUMNS]:
        value = record.get(label)
        if value is None or pd.isna(value) or float(value) <= 0:
            bugs.append(f"{label} must be present and positive.")

    if record["CUR_MKT_CAP(HKD)"] is None or record["CUR_MKT_CAP(HKD)"] <= 0:
        bugs.append("CUR_MKT_CAP(HKD) must be present and positive.")

    # ----- Large-change checks requiring manual review -----
    changes = {
        "China Close": percentage_change(china["Value"], china["Previous_Value"]),
        "HK Close": percentage_change(hk["Value"], hk["Previous_Value"]),
        "FX": percentage_change(fx["Value"], fx["Previous_Value"]),
    }
    if changes["China Close"] is not None and changes["China Close"] > MAX_PRICE_CHANGE:
        review_items.append(f"China Close change {changes['China Close']:.1%}.")
    if changes["HK Close"] is not None and changes["HK Close"] > MAX_PRICE_CHANGE:
        review_items.append(f"HK Close change {changes['HK Close']:.1%}.")
    if changes["FX"] is not None and changes["FX"] > MAX_FX_CHANGE:
        review_items.append(f"FX change {changes['FX']:.1%}.")

    if risk_free["Value"] is not None and risk_free["Previous_Value"] is not None:
        rf_change = abs(risk_free["Value"] - risk_free["Previous_Value"])
        if rf_change > MAX_RISK_FREE_CHANGE_PP:
            review_items.append(f"Risk-free change {rf_change:.2f} percentage points.")

    previous_market_cap = record.get("Previous_CUR_MKT_CAP(HKD)")
    market_cap_change = percentage_change(
        record["CUR_MKT_CAP(HKD)"], previous_market_cap
    )
    if market_cap_change is not None and market_cap_change > MAX_MARKET_CAP_CHANGE:
        review_items.append(f"Market-cap change {market_cap_change:.1%}.")

    if bugs:
        quality_status = "FAIL"
    elif review_items:
        quality_status = "REVIEW_REQUIRED"
    elif warnings_list:
        quality_status = "PASS_WITH_WARNING"
    else:
        quality_status = "PASS"

    ready_for_dtd = record["HK_Open"] and quality_status in {
        "PASS",
        "PASS_WITH_WARNING",
    }

    return {
        "Bug_Flag": bool(bugs),
        "Bug_Message": " | ".join(bugs) if bugs else None,
        "Warning_Flag": bool(warnings_list),
        "Warning_Message": " | ".join(warnings_list) if warnings_list else None,
        "Review_Flag": bool(review_items),
        "Review_Message": " | ".join(review_items) if review_items else None,
        "Quality_Status": quality_status,
        "Ready_For_DTD": ready_for_dtd,
    }


def read_excel_sheet_if_present(path, sheet_name):
    if not path.exists():
        return pd.DataFrame()
    workbook = pd.ExcelFile(path)
    if sheet_name not in workbook.sheet_names:
        return pd.DataFrame()
    return pd.read_excel(path, sheet_name=sheet_name)


def compare_approved_row(existing_row, new_row):
    for column in TEMPORARY_COLUMNS:
        left = existing_row[column]
        right = new_row[column]
        if column == "Date":
            if pd.Timestamp(left).normalize() != pd.Timestamp(right).normalize():
                return False
        elif pd.isna(left) and pd.isna(right):
            continue
        elif not pd.notna(left) or not pd.notna(right):
            return False
        elif abs(float(left) - float(right)) > max(1e-8, abs(float(right)) * 1e-10):
            return False
    return True


def write_temporary_like_vanke(excel_output):
    """Write only approved rows, but preserve the exact vanke.xlsx/Input layout."""
    workbook = load_workbook(CORE_FILE)
    input_sheet = workbook["Input"]

    template_headers = [
        input_sheet.cell(row=1, column=column).value
        for column in range(1, len(TEMPORARY_COLUMNS) + 1)
    ]
    if template_headers != TEMPORARY_COLUMNS:
        raise RuntimeError(
            "ALERT: vanke.xlsx/Input schema changed; Temporary was not overwritten."
        )

    # Capture the reference data-row formatting before removing historical rows.
    template_styles = [
        copy(input_sheet.cell(row=2, column=column)._style)
        for column in range(1, len(TEMPORARY_COLUMNS) + 1)
    ]
    template_row_height = input_sheet.row_dimensions[2].height

    # Temporary is an approved incremental Input file, not a copy of history.
    if input_sheet.max_row > 1:
        input_sheet.delete_rows(2, input_sheet.max_row - 1)

    for worksheet in list(workbook.worksheets):
        if worksheet.title != "Input":
            workbook.remove(worksheet)

    for row_number, values in enumerate(
        excel_output[TEMPORARY_COLUMNS].itertuples(index=False, name=None),
        start=2,
    ):
        for column_number, value in enumerate(values, start=1):
            if pd.isna(value):
                value = None
            elif hasattr(value, "item"):
                value = value.item()
            cell = input_sheet.cell(row=row_number, column=column_number, value=value)
            cell._style = copy(template_styles[column_number - 1])
        input_sheet.row_dimensions[row_number].height = template_row_height

    workbook.save(TEMPORARY_INPUT_FILE)


def update_temporary(record, save_result):
    """Insert only a validated HK trading-day row with the exact Input schema."""
    existing = read_excel_sheet_if_present(TEMPORARY_INPUT_FILE, "Input")
    if not existing.empty:
        existing["Date"] = pd.to_datetime(
            existing["Date"].astype(str).str.replace(r"\.0$", "", regex=True),
            format="%Y%m%d",
        )

    if not record["HK_Open"]:
        return "NOT_ELIGIBLE_HK_CLOSED", False

    if not record["Ready_For_DTD"]:
        return "BLOCKED_BY_DATA_QUALITY", False

    new_row = {
        "Comp_no": record["Comp_no"],
        "Date": record["Date"],
        "CUR_MKT_CAP(HKD)": record["CUR_MKT_CAP(HKD)"],
        "BS_CUR_LIAB(HKD)": record["BS_CUR_LIAB(HKD)"],
        "BS_LT_BORROW(HKD)": record["BS_LT_BORROW(HKD)"],
        "BS_TOT_LIAB2(HKD)": record["BS_TOT_LIAB2(HKD)"],
        "BS_TOT_ASSET(HKD)": record["BS_TOT_ASSET(HKD)"],
        "Risk_Free_Rate": record["Risk_Free_Rate"],
    }

    same_date = (
        existing.loc[existing["Date"] == record["Date"]]
        if not existing.empty
        else pd.DataFrame()
    )
    if not same_date.empty:
        if compare_approved_row(same_date.iloc[-1], new_row):
            return "UNCHANGED_EXISTING_APPROVED_ROW", False
        return "BLOCKED_CONFLICT_WITH_APPROVED_ROW", True

    if save_result:
        updated = (
            pd.DataFrame([new_row])
            if existing.empty
            else pd.concat([existing, pd.DataFrame([new_row])], ignore_index=True)
        )
        updated = updated.sort_values("Date").reset_index(drop=True)
        excel_output = updated[TEMPORARY_COLUMNS].copy()
        excel_output["Date"] = excel_output["Date"].dt.strftime("%Y%m%d").astype(int)
        write_temporary_like_vanke(excel_output)
        return "INSERTED", False

    return "WOULD_INSERT_DRY_RUN", False


def save_daily_datalog(record, attempt_rows):
    daily = read_excel_sheet_if_present(DAILY_DATALOG_FILE, "Daily_Result")
    attempts = read_excel_sheet_if_present(DAILY_DATALOG_FILE, "API_Attempt_Log")

    new_daily = pd.DataFrame([record])
    if not daily.empty:
        daily["Date"] = pd.to_datetime(daily["Date"]).dt.normalize()
    daily = pd.concat([daily, new_daily], ignore_index=True)
    daily["Date"] = pd.to_datetime(daily["Date"]).dt.normalize()
    daily = (
        daily.sort_values("Date")
        .drop_duplicates(subset=["Date"], keep="last")
        .reset_index(drop=True)
    )

    if attempt_rows:
        attempts = pd.concat([attempts, pd.DataFrame(attempt_rows)], ignore_index=True)

    with pd.ExcelWriter(DAILY_DATALOG_FILE, engine="openpyxl") as writer:
        daily.to_excel(writer, sheet_name="Daily_Result", index=False)
        attempts.to_excel(writer, sheet_name="API_Attempt_Log", index=False)


# ============================================================
# MODULE 4 - PROCESS ONE INPUT DATE
# ============================================================


def process_daily_data(input_date, save_result=True):
    run_id = uuid.uuid4().hex
    run_time = datetime.now()
    data_date = pd.to_datetime(str(input_date), format="%Y%m%d").normalize()

    calendar = pd.read_excel(CALENDAR_FILE, sheet_name="Daily Calendar")
    calendar["Date"] = pd.to_datetime(calendar["Date"]).dt.normalize()
    calendar_row = calendar.loc[calendar["Date"] == data_date]
    if calendar_row.empty:
        raise ValueError(f"{data_date.date()} is missing from Daily Calendar.")

    china_open = bool(calendar_row.iloc[0]["China_Open"])
    hk_open = bool(calendar_row.iloc[0]["HK_Open"])
    if china_open and hk_open:
        market_case = "BOTH_OPEN"
    elif china_open:
        market_case = "CHINA_ONLY"
    elif hk_open:
        market_case = "HK_ONLY"
    else:
        market_case = "BOTH_CLOSED"

    china = get_online_close(
        "CHINA_CLOSE", CHINA_TICKER, data_date, require_exact=china_open
    )
    hk = get_online_close("HK_CLOSE", HK_TICKER, data_date, require_exact=hk_open)
    fx = get_online_close("FX", FX_TICKER, data_date, require_exact=hk_open)
    risk_free = get_hkma_risk_free(
        data_date, require_exact=hk_open, update_cache=save_result
    )
    sources = {
        "CHINA_CLOSE": china,
        "HK_CLOSE": hk,
        "FX": fx,
        "RISK_FREE": risk_free,
    }

    capital = pd.read_excel(ISSUED_CAPITAL_FILE)
    capital["Time"] = pd.to_datetime(capital["Time"]).dt.normalize()
    capital = capital.loc[capital["Time"] <= data_date].sort_values("Time")
    if capital.empty:
        raise ValueError("No issued-capital data are available as of the input date.")
    capital_row = capital.iloc[-1]

    full_core = pd.read_excel(CORE_FILE, sheet_name="Input")
    full_core["Date"] = pd.to_datetime(full_core["Date"].astype(str), format="%Y%m%d")
    core = full_core.loc[full_core["Date"] <= data_date].sort_values("Date")
    if core.empty:
        raise ValueError("No balance-sheet data are available as of the input date.")
    core_row = core.iloc[-1]
    use_company_datalog_financials = all(
        column in capital_row.index and pd.notna(capital_row[column])
        for column in BALANCE_SHEET_COLUMNS
    )
    balance_sheet_row = capital_row if use_company_datalog_financials else core_row
    balance_sheet_source_date = (
        pd.Timestamp(capital_row["Time"]).normalize()
        if use_company_datalog_financials
        else pd.Timestamp(core_row["Date"]).normalize()
    )

    china_shares = float(capital_row["China Stock(A)"])
    hk_shares = float(capital_row["HongKong Stock(H)"])
    fx_value = fx["Value"]
    china_value = china["Value"]
    hk_value = hk["Value"]

    if None not in (china_value, hk_value, fx_value):
        china_market_cap = china_value * china_shares * fx_value / 1_000_000
        hk_market_cap = hk_value * hk_shares / 1_000_000
        total_market_cap = china_market_cap + hk_market_cap
    else:
        china_market_cap = None
        hk_market_cap = None
        total_market_cap = None

    # QC must compare with the latest available observation, including a prior
    # validated Temporary row.  Comparing every new day only with the final
    # confirmed date would hide or exaggerate day-on-day moves.
    confirmed_previous = full_core.loc[
        full_core["Date"] < data_date, ["Date", "CUR_MKT_CAP(HKD)"]
    ].copy()
    confirmed_previous["_source_priority"] = 2
    previous_frames = [confirmed_previous]
    previous_temporary = read_excel_sheet_if_present(TEMPORARY_INPUT_FILE, "Input")
    if not previous_temporary.empty:
        previous_temporary["Date"] = pd.to_datetime(
            previous_temporary["Date"].astype(str).str.replace(r"\.0$", "", regex=True),
            format="%Y%m%d",
        )
        temporary_previous = previous_temporary.loc[
            previous_temporary["Date"] < data_date,
            ["Date", "CUR_MKT_CAP(HKD)"],
        ].copy()
        temporary_previous["_source_priority"] = 1
        previous_frames.append(temporary_previous)
    previous_core = (
        pd.concat(previous_frames, ignore_index=True)
        .sort_values(["Date", "_source_priority"])
        .drop_duplicates("Date", keep="last")
        .sort_values("Date")
    )
    previous_market_cap = (
        float(previous_core.iloc[-1]["CUR_MKT_CAP(HKD)"])
        if not previous_core.empty
        else None
    )

    record = {
        "Run_ID": run_id,
        "Run_Time": run_time,
        "Comp_no": int(core_row["Comp_no"]),
        "Date": data_date,
        "China_Open": china_open,
        "HK_Open": hk_open,
        "Market_Case": market_case,
        "China_Close": china_value if china_open else None,
        "China_Close_Used": china_value,
        "China_Close_Source_Date": china["Source_Date"],
        "China_API_Status": china["Status"],
        "China_API_Attempts": len(china["Attempts"]),
        "HK_Close": hk_value if hk_open else None,
        "HK_Close_Used": hk_value,
        "HK_Close_Source_Date": hk["Source_Date"],
        "HK_API_Status": hk["Status"],
        "HK_API_Attempts": len(hk["Attempts"]),
        "FX_Rate": fx_value,
        "FX_Source_Date": fx["Source_Date"],
        "FX_API_Status": fx["Status"],
        "FX_API_Attempts": len(fx["Attempts"]),
        "Risk_Free_Rate": risk_free["Value"],
        "Risk_Free_Source_Date": risk_free["Source_Date"],
        "Risk_Free_API_Status": risk_free["Status"],
        "Risk_Free_API_Attempts": len(risk_free["Attempts"]),
        "China_Shares": china_shares,
        "HK_Shares": hk_shares,
        "Capital_Source_Date": pd.Timestamp(capital_row["Time"]).normalize(),
        "China_Market_Cap_HKD_mn": china_market_cap,
        "HK_Market_Cap_HKD_mn": hk_market_cap,
        "CUR_MKT_CAP(HKD)": total_market_cap,
        "Previous_CUR_MKT_CAP(HKD)": previous_market_cap,
        "Balance_Sheet_Source_Date": balance_sheet_source_date,
    }
    for column in BALANCE_SHEET_COLUMNS:
        record[column] = float(balance_sheet_row[column])

    quality = check_daily_data(record, sources)
    record.update(quality)

    temporary_action, conflict = update_temporary(record, save_result)
    if conflict:
        record["Bug_Flag"] = True
        conflict_message = (
            "New values conflict with an existing approved Temporary row."
        )
        record["Bug_Message"] = " | ".join(
            item for item in [record["Bug_Message"], conflict_message] if item
        )
        record["Quality_Status"] = "FAIL"
        record["Ready_For_DTD"] = False
    record["Temporary_Action"] = temporary_action

    attempt_rows = []
    for source in sources.values():
        for attempt in source["Attempts"]:
            attempt_rows.append(
                {
                    "Run_ID": run_id,
                    "Input_Date": data_date,
                    **attempt,
                }
            )

    if save_result:
        save_daily_datalog(record, attempt_rows)

    result = pd.DataFrame([record])

    if record["Bug_Flag"] or record["Review_Flag"]:
        save_note = (
            "Daily Datalog was saved."
            if save_result
            else "Dry run; no file was changed."
        )
        raise RuntimeError(
            "ALERT: data did not pass the Temporary gate. "
            f"{save_note} Temporary was not updated. "
            f"Bug={record['Bug_Message']} Review={record['Review_Message']}"
        )

    return result
