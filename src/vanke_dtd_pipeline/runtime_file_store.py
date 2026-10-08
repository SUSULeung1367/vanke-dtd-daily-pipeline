"""Excel persistence for generated workspaces only.

The tracked files under ``data/`` are never written by these functions.
"""

from __future__ import annotations

from copy import copy

import pandas as pd
from openpyxl import load_workbook

from .pipeline_config import INPUT_COLUMNS
from .runtime_workspace import current_workspace


def read_excel_sheet_if_present(path, sheet_name):
    if not path.exists():
        return pd.DataFrame()
    with pd.ExcelFile(path) as workbook:
        if sheet_name not in workbook.sheet_names:
            return pd.DataFrame()
    return pd.read_excel(path, sheet_name=sheet_name)


def compare_approved_row(existing_row, new_row):
    for column in INPUT_COLUMNS:
        left, right = existing_row[column], new_row[column]
        if column == "Date":
            if pd.Timestamp(left).normalize() != pd.Timestamp(right).normalize():
                return False
        elif pd.isna(left) and pd.isna(right):
            continue
        elif not pd.notna(left) or not pd.notna(right):
            return False
        elif abs(float(left) - float(right)) > max(1e-8, abs(float(right)) * 1e-10):
            return False
    return True


def write_temporary_input(excel_output):
    """Write validated incremental rows using the confirmed Input schema/style."""
    paths = current_workspace()
    workbook = load_workbook(paths.confirmed_history)
    input_sheet = workbook["Input"]
    headers = [input_sheet.cell(1, column).value for column in range(1, len(INPUT_COLUMNS) + 1)]
    if headers != INPUT_COLUMNS:
        raise RuntimeError("Confirmed Input schema changed; pending-review Input was not overwritten.")
    template_styles = [copy(input_sheet.cell(2, column)._style) for column in range(1, len(INPUT_COLUMNS) + 1)]
    template_row_height = input_sheet.row_dimensions[2].height
    if input_sheet.max_row > 1:
        input_sheet.delete_rows(2, input_sheet.max_row - 1)
    for worksheet in list(workbook.worksheets):
        if worksheet.title != "Input":
            workbook.remove(worksheet)
    for row_number, values in enumerate(excel_output[INPUT_COLUMNS].itertuples(index=False, name=None), start=2):
        for column_number, value in enumerate(values, start=1):
            clean_value = None if pd.isna(value) else value.item() if hasattr(value, "item") else value
            cell = input_sheet.cell(row_number, column_number, value=clean_value)
            cell._style = copy(template_styles[column_number - 1])
        input_sheet.row_dimensions[row_number].height = template_row_height
    workbook.save(paths.pending_review_input)


def update_temporary(record, save_result):
    """Insert one QC-approved HK-open row; never overwrite a prior temporary row."""
    paths = current_workspace()
    existing = read_excel_sheet_if_present(paths.pending_review_input, "Input")
    if not existing.empty:
        existing["Date"] = pd.to_datetime(existing["Date"].astype(str).str.replace(r"\.0$", "", regex=True), format="%Y%m%d")
    if not record["HK_Open"]:
        return "NOT_ELIGIBLE_HK_CLOSED", False
    if not record["Ready_For_DTD"]:
        return "BLOCKED_BY_DATA_QUALITY", False
    new_row = {column: record[column] for column in INPUT_COLUMNS}
    same_date = existing.loc[existing["Date"] == record["Date"]] if not existing.empty else pd.DataFrame()
    if not same_date.empty:
        if compare_approved_row(same_date.iloc[-1], new_row):
            return "UNCHANGED_EXISTING_APPROVED_ROW", False
        return "BLOCKED_CONFLICT_WITH_APPROVED_ROW", True
    if not save_result:
        return "WOULD_INSERT_DRY_RUN", False
    updated = pd.DataFrame([new_row]) if existing.empty else pd.concat([existing, pd.DataFrame([new_row])], ignore_index=True)
    updated = updated.sort_values("Date").reset_index(drop=True)
    output = updated[INPUT_COLUMNS].copy()
    output["Date"] = output["Date"].dt.strftime("%Y%m%d").astype(int)
    write_temporary_input(output)
    return "INSERTED", False


def save_daily_datalog(record, attempt_rows):
    """Append/update the generated run audit file, keyed by data date."""
    path = current_workspace().market_data_audit
    daily = read_excel_sheet_if_present(path, "Daily_Result")
    attempts = read_excel_sheet_if_present(path, "API_Attempt_Log")
    daily = pd.concat([daily, pd.DataFrame([record])], ignore_index=True)
    daily["Date"] = pd.to_datetime(daily["Date"]).dt.normalize()
    daily = daily.sort_values("Date").drop_duplicates("Date", keep="last").reset_index(drop=True)
    if attempt_rows:
        attempts = pd.concat([attempts, pd.DataFrame(attempt_rows)], ignore_index=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        daily.to_excel(writer, sheet_name="Daily_Result", index=False)
        attempts.to_excel(writer, sheet_name="API_Attempt_Log", index=False)
