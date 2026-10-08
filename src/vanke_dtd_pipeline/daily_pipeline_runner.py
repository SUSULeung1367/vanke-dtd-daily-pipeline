"""Standard runner shared by the live daily run, basic test and demos.

This is the only end-to-end implementation. The basic test and notebooks call
this module; they do not contain a second DTD calculation.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional
import shutil

import pandas as pd
from openpyxl import load_workbook

from . import daily_input_builder
from .pipeline_config import DEFAULT_COMPANY, INPUT_COLUMNS
from .daily_dtd_calculator import (
    DailyProcessingResult,
    combine_input_history,
    process_daily_dtd,
)
from .project_data_paths import repository_data, repository_root, validate_repository_data
from .runtime_workspace import WorkspacePaths


@dataclass(frozen=True)
class PipelineResult:
    date: str
    market_case: str
    quality_status: str
    pending_review_input_action: str
    dtd_status: str
    dtd: Optional[float]
    wrote_dtd_output: bool
    message: str

    def as_dict(self) -> dict:
        return asdict(self)


@dataclass(frozen=True)
class RangeRunResult:
    start_date: str
    end_date: str
    mode: str
    workspace: Path
    daily_results: pd.DataFrame
    pending_review_input: pd.DataFrame
    pending_review_dtd_output: pd.DataFrame
    updated_clean_input: pd.DataFrame


def run_one_date(
    input_date,
    *,
    project_dir=None,
    company=None,
    save_result=True,
) -> PipelineResult:
    """Run ingestion/QC and then the calendar-controlled DTD calculation.

    ``save_result=True`` is required for the two file-based stages to hand off
    the newly validated Input row.  Use ``run_prepared_date`` for a safe test
    against an already prepared temporary Input file.
    """
    if not save_result:
        raise ValueError(
            "End-to-end dry run cannot hand off a new row between file-based "
            "stages. Use run_prepared_date() for a no-production-file test."
        )

    root = Path(project_dir or repository_root()).resolve()
    # A repository root holds tracked inputs under data/; a workspace holds
    # copied filenames at its own root. Support both without risking a write to
    # the tracked baseline.
    if (root / "data").is_dir():
        root = prepare_runtime_workspace(root, reset=False)
    paths = WorkspacePaths(root)
    daily_input_builder.configure_runtime_workspace(root)
    daily = daily_input_builder.process_daily_data(input_date, save_result=True).iloc[0]

    dtd_result = process_daily_dtd(
        input_date,
        confirmed_file=paths.confirmed_history,
        temporary_input_file=paths.pending_review_input,
        temporary_output_file=paths.pending_review_output,
        calendar_file=paths.calendar,
        company=company,
    )
    return _combine_results(daily, dtd_result)


def run_prepared_date(
    input_date,
    *,
    confirmed_file,
    temporary_input_file,
    temporary_output_file,
    calendar_file,
    company=None,
    write_output=True,
) -> DailyProcessingResult:
    """Test the hand-off contract using a prepared, QC-approved Input file."""
    return process_daily_dtd(
        input_date,
        confirmed_file=confirmed_file,
        temporary_input_file=temporary_input_file,
        temporary_output_file=temporary_output_file,
        calendar_file=calendar_file,
        company=company,
        write_output=write_output,
    )


def prepare_runtime_workspace(project_dir=None, workspace_dir=None, reset=True) -> Path:
    """Create a runtime workspace without changing tracked standard inputs."""
    source = Path(project_dir or repository_root()).resolve()
    data = repository_data(source)
    validate_repository_data(data)
    workspace = Path(workspace_dir or source / "runtime" / "pipeline_workspace").resolve()
    workspace.mkdir(parents=True, exist_ok=True)
    paths = WorkspacePaths(workspace)

    required = {
        data.confirmed_history: paths.confirmed_history,
        data.trading_calendar: paths.calendar,
        data.company_data: paths.company_data,
        data.risk_free_rate_cache: paths.risk_free_rate_cache,
    }

    if reset or not paths.confirmed_history.exists():
        for source_path, workspace_path in required.items():
            shutil.copy2(source_path, workspace_path)
        for path in (
            paths.pending_review_input,
            paths.pending_review_output,
            paths.market_data_audit,
        ):
            if path.exists():
                path.unlink()
        _create_empty_pending_review_input(paths)
    return workspace


def run_date_range(
    start_date,
    end_date,
    *,
    project_dir=None,
    workspace_dir=None,
    mode="BASIC_TEST",
    reset=True,
    company=DEFAULT_COMPANY,
) -> RangeRunResult:
    """Process an inclusive range with the standard daily pipeline.

    BASIC_TEST reads a fixed saved market-data fixture to verify the local
    installation. LIVE uses online sources. Both modes call the same daily
    input and DTD calculation code, and write only inside a runtime workspace.
    """
    source = Path(project_dir or repository_root()).resolve()
    workspace = prepare_runtime_workspace(source, workspace_dir, reset=reset)
    paths = WorkspacePaths(workspace)
    mode = str(mode).strip().upper()
    if mode not in {"BASIC_TEST", "LIVE"}:
        raise ValueError("mode must be BASIC_TEST or LIVE")

    start = _parse_date(start_date)
    end = _parse_date(end_date)
    if start > end:
        raise ValueError("START_DATE must not be later than END_DATE")

    calendar = pd.read_excel(
        paths.calendar, sheet_name="Daily Calendar"
    )
    calendar["Date"] = pd.to_datetime(calendar["Date"]).dt.normalize()
    selected = calendar.loc[calendar["Date"].between(start, end)].copy()
    expected_days = (end - start).days + 1
    if len(selected) != expected_days:
        raise ValueError(
            "The selected range is not fully covered by the trading calendar"
        )

    daily_input_builder.configure_runtime_workspace(workspace)
    basic_test_data = _load_basic_test_market_data(source) if mode == "BASIC_TEST" else pd.DataFrame()
    summaries = []
    saved_daily_rows = []

    for day in pd.date_range(start, end, freq="D"):
        day_text = day.strftime("%Y%m%d")
        calendar_row = selected.loc[selected["Date"] == day].iloc[0]
        hk_open = bool(calendar_row["HK_Open"])
        china_open = bool(calendar_row["China_Open"])

        try:
            if mode == "LIVE":
                daily = daily_input_builder.process_daily_data(day_text, save_result=True).iloc[0]
            else:
                daily = _basic_test_daily_row(day, calendar_row, basic_test_data)
                if "Pending_Review_Input_Action" not in daily:
                    daily["Pending_Review_Input_Action"] = daily.get("Temporary_Action")
                if hk_open:
                    action, conflict = daily_input_builder.update_temporary(daily, True)
                    if conflict:
                        raise RuntimeError("Basic-test row conflicts with pending-review Input")
                    daily["Pending_Review_Input_Action"] = action
                saved_daily_rows.append(dict(daily))

            dtd = process_daily_dtd(
                day_text,
                confirmed_file=paths.confirmed_history,
                temporary_input_file=paths.pending_review_input,
                temporary_output_file=paths.pending_review_output,
                calendar_file=paths.calendar,
                company=company,
            )
            summaries.append(
                {
                    "Date": day.strftime("%Y-%m-%d"),
                    "China_Open": china_open,
                    "HK_Open": hk_open,
                    "Market_Case": daily.get("Market_Case"),
                    "Input_QC": daily.get("Quality_Status"),
                    "Pending_Review_Input_Action": daily.get("Pending_Review_Input_Action"),
                    "DTD_Status": dtd.status,
                    "DTD": dtd.dtd,
                    "Message": dtd.message,
                }
            )
        except Exception as error:
            summaries.append(
                {
                    "Date": day.strftime("%Y-%m-%d"),
                    "China_Open": china_open,
                    "HK_Open": hk_open,
                    "Market_Case": _market_case(china_open, hk_open),
                    "Input_QC": "BLOCKED",
                    "Pending_Review_Input_Action": "NOT_WRITTEN",
                    "DTD_Status": type(error).__name__,
                    "DTD": None,
                    "Message": str(error),
                }
            )

    if mode == "BASIC_TEST":
        _write_basic_test_audit(paths.market_data_audit, saved_daily_rows)

    temporary_input = pd.read_excel(
        paths.pending_review_input, sheet_name="Input"
    )
    temporary_dtd_path = paths.pending_review_output
    temporary_dtd = (
        pd.read_excel(temporary_dtd_path, sheet_name="Output")
        if temporary_dtd_path.exists()
        else pd.DataFrame(columns=["Comp_no", "Date", "DTD"])
    )
    confirmed_input = pd.read_excel(paths.confirmed_history, sheet_name="Input")
    updated_clean_input = combine_input_history(
        confirmed_input, temporary_input, company=company
    )[INPUT_COLUMNS].copy()
    updated_clean_input["Date"] = (
        pd.to_datetime(updated_clean_input["Date"]).dt.strftime("%Y%m%d").astype(int)
    )
    return RangeRunResult(
        start_date=start.strftime("%Y-%m-%d"),
        end_date=end.strftime("%Y-%m-%d"),
        mode=mode,
        workspace=workspace,
        daily_results=pd.DataFrame(summaries),
        pending_review_input=temporary_input,
        pending_review_dtd_output=temporary_dtd,
        updated_clean_input=updated_clean_input,
    )


def _combine_results(daily: pd.Series, dtd: DailyProcessingResult) -> PipelineResult:
    return PipelineResult(
        date=pd.Timestamp(daily["Date"]).strftime("%Y-%m-%d"),
        market_case=str(daily["Market_Case"]),
        quality_status=str(daily["Quality_Status"]),
        pending_review_input_action=str(daily["Pending_Review_Input_Action"]),
        dtd_status=dtd.status,
        dtd=dtd.dtd,
        wrote_dtd_output=dtd.wrote_output,
        message=dtd.message,
    )


def _parse_date(value) -> pd.Timestamp:
    text = str(value).strip().removesuffix(".0")
    value = pd.to_datetime(
        text, format="%Y%m%d" if len(text) == 8 and text.isdigit() else None
    )
    return pd.Timestamp(value).normalize()


def _market_case(china_open: bool, hk_open: bool) -> str:
    if china_open and hk_open:
        return "BOTH_OPEN"
    if china_open:
        return "CHINA_ONLY"
    if hk_open:
        return "HK_ONLY"
    return "BOTH_CLOSED"


def _create_empty_pending_review_input(paths: WorkspacePaths) -> None:
    workbook = load_workbook(paths.confirmed_history)
    sheet = workbook["Input"]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for other in list(workbook.worksheets):
        if other.title != "Input":
            workbook.remove(other)
    workbook.save(paths.pending_review_input)


def _load_basic_test_market_data(source: Path) -> pd.DataFrame:
    path = repository_data(source).basic_test_market_data
    if not path.exists():
        raise FileNotFoundError("Basic-test market-data fixture is missing")
    fixture = pd.read_excel(path, sheet_name="Daily_Result")
    fixture["Date"] = pd.to_datetime(fixture["Date"]).dt.normalize()
    return fixture.sort_values("Date").drop_duplicates("Date", keep="last")


def _basic_test_daily_row(day, calendar_row, basic_test_data: pd.DataFrame) -> dict:
    source_row = basic_test_data.loc[basic_test_data["Date"] == day]
    china_open = bool(calendar_row["China_Open"])
    hk_open = bool(calendar_row["HK_Open"])
    if not source_row.empty:
        return source_row.iloc[-1].to_dict()
    if hk_open:
        raise ValueError(
            f"No basic-test market-data row exists for {day.date()}; use LIVE mode or choose a supported test date"
        )
    return {
        "Date": day,
        "China_Open": china_open,
        "HK_Open": hk_open,
        "Market_Case": _market_case(china_open, hk_open),
        "Quality_Status": "PASS_EXPECTED_HK_CLOSED",
        "Ready_For_DTD": False,
        "Bug_Flag": False,
        "Review_Flag": False,
        "Temporary_Action": "NOT_ELIGIBLE_HK_CLOSED",
    }


def _write_basic_test_audit(audit_path: Path, rows) -> None:
    daily = pd.DataFrame(rows)
    with pd.ExcelWriter(
        audit_path, engine="openpyxl"
    ) as writer:
        daily.to_excel(writer, sheet_name="Daily_Result", index=False)
        pd.DataFrame(
            columns=[
                "Run_ID",
                "Input_Date",
                "API",
                "Attempt",
                "Attempt_Status",
                "Attempt_Time",
                "Error_Type",
                "Error_Message",
            ]
        ).to_excel(writer, sheet_name="API_Attempt_Log", index=False)
