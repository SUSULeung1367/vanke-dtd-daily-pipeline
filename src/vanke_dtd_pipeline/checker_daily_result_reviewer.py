"""Checker review and append-only release for the Vanke DTD demonstration."""

from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
import os
import uuid

import pandas as pd
from openpyxl import load_workbook

from .runtime_workspace import WorkspacePaths


INPUT_COLUMNS = [
    "Comp_no",
    "Date",
    "CUR_MKT_CAP(HKD)",
    "BS_CUR_LIAB(HKD)",
    "BS_LT_BORROW(HKD)",
    "BS_TOT_LIAB2(HKD)",
    "BS_TOT_ASSET(HKD)",
    "Risk_Free_Rate",
]
DTD_COLUMNS = ["Comp_no", "Date", "DTD"]


def build_checker_table(workspace_dir) -> pd.DataFrame:
    """Return one review table containing pending-review Input, DTD and QC."""
    paths = WorkspacePaths(Path(workspace_dir).resolve())
    inputs = pd.read_excel(paths.pending_review_input, sheet_name="Input")
    dtd_path = paths.pending_review_output
    dtd = (
        pd.read_excel(dtd_path, sheet_name="Output")
        if dtd_path.exists()
        else pd.DataFrame(columns=DTD_COLUMNS)
    )
    log_path = paths.market_data_audit
    daily = (
        pd.read_excel(log_path, sheet_name="Daily_Result")
        if log_path.exists()
        else pd.DataFrame()
    )

    inputs["_Date"] = _parse_dates(inputs["Date"])
    dtd["_Date"] = (
        _parse_dates(dtd["Date"])
        if not dtd.empty
        else pd.Series(dtype="datetime64[ns]")
    )
    review = inputs.merge(
        dtd[["Comp_no", "_Date", "DTD"]], on=["Comp_no", "_Date"], how="left"
    )
    if not daily.empty:
        daily["_Date"] = _parse_dates(daily["Date"])
        qc_columns = [
            "_Date",
            "Market_Case",
            "Quality_Status",
            "Bug_Flag",
            "Review_Flag",
            "Ready_For_DTD",
            "Warning_Message",
        ]
        qc_columns = [column for column in qc_columns if column in daily.columns]
        review = review.merge(
            daily[qc_columns]
            .sort_values("_Date")
            .drop_duplicates("_Date", keep="last"),
            on="_Date",
            how="left",
        )

    review["Automated_QC"] = review.apply(_qc_status, axis=1)
    review["Marker_Required"] = (
        review.get("Review_Flag", False).fillna(False).astype(bool)
    )
    review["Production_Status"] = _production_status(paths, review)
    review["Date"] = review["_Date"].dt.strftime("%Y%m%d").astype(int)
    display_columns = [
        *INPUT_COLUMNS,
        "DTD",
        "Market_Case",
        "Quality_Status",
        "Automated_QC",
        "Marker_Required",
        "Production_Status",
        "Warning_Message",
    ]
    return review[
        [column for column in display_columns if column in review.columns]
    ].sort_values("Date")


def confirm_dates(
    workspace_dir, dates, checker_name, decision="APPROVE"
) -> pd.DataFrame:
    """Apply a Checker decision and atomically append approved normal rows."""
    workspace = Path(workspace_dir).resolve()
    decision = str(decision).strip().upper()
    if decision not in {"APPROVE", "REJECT"}:
        raise ValueError("decision must be APPROVE or REJECT")
    checker_name = str(checker_name).strip()
    if not checker_name:
        raise ValueError("checker_name is required")

    review = build_checker_table(workspace)
    selected_dates = {_date_key(value) for value in dates}
    selected = review.loc[review["Date"].astype(str).isin(selected_dates)].copy()
    if selected.empty:
        raise ValueError("No selected dates were found in the Checker table")
    if len(selected) != len(selected_dates):
        found = set(selected["Date"].astype(str))
        raise ValueError(
            f"Dates not found in the Checker table: {sorted(selected_dates - found)}"
        )

    if decision == "APPROVE":
        blocked = selected.loc[
            (selected["Automated_QC"] != "PASS")
            | selected["DTD"].isna()
            | selected["Marker_Required"]
            | (selected["Production_Status"] == "RELEASED")
        ]
        if not blocked.empty:
            raise ValueError(
                "Normal Checker approval is blocked for dates: "
                + ", ".join(blocked["Date"].astype(str))
            )

    _write_release(workspace, selected, checker_name, decision)
    return build_checker_table(workspace)


def _write_release(
    workspace: Path, selected: pd.DataFrame, checker_name: str, decision: str
) -> None:
    paths = WorkspacePaths(workspace)
    confirmed_path = paths.confirmed_history
    workbook = load_workbook(confirmed_path)
    input_sheet = workbook["Input"]
    output_sheet = workbook["Output"]
    log_sheet = (
        workbook["Confirmation_Log"]
        if "Confirmation_Log" in workbook.sheetnames
        else workbook.create_sheet("Confirmation_Log")
    )
    log_headers = [
        "Confirmation_ID",
        "Confirmation_Time",
        "Data_Date",
        "Checker",
        "Checker_Decision",
        "Automated_QC",
        "Marker_Required",
        "Release_Status",
    ]
    if log_sheet.max_row == 1 and log_sheet.cell(1, 1).value is None:
        for column, header in enumerate(log_headers, start=1):
            log_sheet.cell(1, column, header)

    existing_input = {
        _date_key(input_sheet.cell(row, 2).value)
        for row in range(2, input_sheet.max_row + 1)
    }
    existing_output = {
        _date_key(output_sheet.cell(row, 2).value)
        for row in range(2, output_sheet.max_row + 1)
    }

    for _, row in selected.sort_values("Date").iterrows():
        date_key = str(int(row["Date"]))
        release_status = "REJECTED"
        if decision == "APPROVE":
            if date_key in existing_input or date_key in existing_output:
                raise ValueError(
                    f"Confirmed history already contains {date_key}; overwrite is not allowed"
                )
            _append_styled_row(input_sheet, [row[column] for column in INPUT_COLUMNS])
            _append_styled_row(output_sheet, [row[column] for column in DTD_COLUMNS])
            existing_input.add(date_key)
            existing_output.add(date_key)
            release_status = "RELEASED"
        log_sheet.append(
            [
                uuid.uuid4().hex,
                datetime.now(),
                int(date_key),
                checker_name,
                decision,
                row["Automated_QC"],
                bool(row["Marker_Required"]),
                release_status,
            ]
        )

    temp_path = workspace / f".confirmed_history_update_{uuid.uuid4().hex}.xlsx"
    try:
        workbook.save(temp_path)
        os.replace(temp_path, confirmed_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _append_styled_row(sheet, values) -> None:
    target_row = sheet.max_row + 1
    template_row = max(2, sheet.max_row)
    for column, value in enumerate(values, start=1):
        if pd.isna(value):
            value = None
        elif hasattr(value, "item"):
            value = value.item()
        cell = sheet.cell(target_row, column, value=value)
        if sheet.max_row >= 2:
            cell._style = copy(sheet.cell(template_row, column)._style)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[template_row].height


def _qc_status(row) -> str:
    quality = str(row.get("Quality_Status", ""))
    ready = bool(row.get("Ready_For_DTD", False))
    bug = bool(row.get("Bug_Flag", False))
    review = bool(row.get("Review_Flag", False))
    return (
        "PASS"
        if quality in {"PASS", "PASS_WITH_WARNING"}
        and ready
        and not bug
        and not review
        and pd.notna(row.get("DTD"))
        else "BLOCKED"
    )


def _production_status(paths: WorkspacePaths, review: pd.DataFrame) -> pd.Series:
    confirmed = pd.read_excel(paths.confirmed_history, sheet_name="Output")
    confirmed_dates = set(_parse_dates(confirmed["Date"]).dt.strftime("%Y%m%d"))
    return (
        review["_Date"]
        .dt.strftime("%Y%m%d")
        .map(
            lambda value: "RELEASED" if value in confirmed_dates else "PENDING_CHECKER"
        )
    )


def _parse_dates(series: pd.Series) -> pd.Series:
    text = series.astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    compact = text.str.fullmatch(r"\d{8}")
    parsed = pd.to_datetime(text, errors="coerce")
    parsed.loc[compact] = pd.to_datetime(
        text.loc[compact], format="%Y%m%d", errors="coerce"
    )
    return parsed.dt.normalize()


def _date_key(value) -> str:
    text = str(value).strip().removesuffix(".0")
    parsed = pd.to_datetime(
        text, format="%Y%m%d" if len(text) == 8 and text.isdigit() else None
    )
    return pd.Timestamp(parsed).strftime("%Y%m%d")
