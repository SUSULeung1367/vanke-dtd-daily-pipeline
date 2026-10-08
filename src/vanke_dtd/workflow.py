"""One public entry point for the Vanke daily-input and DTD modules.

The ingestion/QC stage remains separate from the DTD calculation stage.  This
module only coordinates them and returns both results to the caller.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Optional
import shutil

import pandas as pd
from openpyxl import load_workbook

from . import ingestion
from .constants import INPUT_COLUMNS
from .dtd import (
    DailyProcessingResult,
    combine_input_history,
    process_daily_dtd,
)
from .repository import repository_data, repository_root, validate_repository_data


@dataclass(frozen=True)
class PipelineResult:
    date: str
    market_case: str
    quality_status: str
    temporary_action: str
    dtd_status: str
    dtd: Optional[float]
    wrote_dtd_output: bool
    message: str

    def as_dict(self) -> dict:
        return asdict(self)


@dataclass(frozen=True)
class RangeRunResult:
    start_date: str
    end_date: str
    mode: str
    workspace: Path
    daily_results: pd.DataFrame
    temporary_input: pd.DataFrame
    temporary_dtd: pd.DataFrame
    updated_clean_input: pd.DataFrame


def run_one_date(
    input_date,
    *,
    project_dir=None,
    company=None,
    save_result=True,
) -> PipelineResult:
    """Run ingestion/QC and then the calendar-controlled DTD calculation.

    ``save_result=True`` is required for the two file-based stages to hand off
    the newly validated Input row.  Use ``run_prepared_date`` for a safe test
    against an already prepared temporary Input file.
    """
    if not save_result:
        raise ValueError(
            "End-to-end dry run cannot hand off a new row between file-based "
            "stages. Use run_prepared_date() for a no-production-file test."
        )

    root = Path(project_dir or repository_root()).resolve()
    # A repository root holds tracked inputs under data/; a workspace holds
    # copied filenames at its own root. Support both without risking a write to
    # the tracked baseline.
    if (root / "data").is_dir():
        root = prepare_demo_workspace(root, reset=False)
    ingestion.configure_project(root)
    daily = ingestion.process_daily_data(input_date, save_result=True).iloc[0]

    dtd_result = process_daily_dtd(
        input_date,
        confirmed_file=root / "vanke.xlsx",
        temporary_input_file=root / "vanke_dtd_temporary_data.xlsx",
        temporary_output_file=root / "temporary_output.xlsx",
        calendar_file=root / "China_HK_Trading_Calendar.xlsx",
        company=company,
    )
    return _combine_results(daily, dtd_result)


def run_prepared_date(
    input_date,
    *,
    confirmed_file,
    temporary_input_file,
    temporary_output_file,
    calendar_file,
    company=None,
    write_output=True,
) -> DailyProcessingResult:
    """Test the hand-off contract using a prepared, QC-approved Input file."""
    return process_daily_dtd(
        input_date,
        confirmed_file=confirmed_file,
        temporary_input_file=temporary_input_file,
        temporary_output_file=temporary_output_file,
        calendar_file=calendar_file,
        company=company,
        write_output=write_output,
    )


def prepare_demo_workspace(project_dir=None, workspace_dir=None, reset=True) -> Path:
    """Create an isolated workspace without changing tracked source inputs."""
    source = Path(project_dir or repository_root()).resolve()
    data = repository_data(source)
    validate_repository_data(data)
    workspace = Path(workspace_dir or source / "runtime" / "demo_workspace").resolve()
    workspace.mkdir(parents=True, exist_ok=True)

    required = {
        data.baseline_vanke: workspace / "vanke.xlsx",
        data.trading_calendar: workspace / "China_HK_Trading_Calendar.xlsx",
        data.issued_capital_datalog: workspace / "Vanke Issued Capital DataLog.xlsx",
        data.risk_free_cache: workspace / "HKMA_Risk_Free_Daily.xlsx",
    }

    if reset or not (workspace / "vanke.xlsx").exists():
        for source_path, workspace_path in required.items():
            shutil.copy2(source_path, workspace_path)
        for name in (
            "vanke_dtd_temporary_data.xlsx",
            "temporary_output.xlsx",
            "Vanke_Daily_Datalog.xlsx",
        ):
            path = workspace / name
            if path.exists():
                path.unlink()
        _create_empty_temporary_input(workspace)
    return workspace


def run_date_range(
    start_date,
    end_date,
    *,
    project_dir=None,
    workspace_dir=None,
    mode="REPLAY",
    reset=True,
    company=5338,
) -> RangeRunResult:
    """Process every calendar date in an inclusive teacher-selected range.

    REPLAY uses the saved source/QC observations for a deterministic demo.
    LIVE uses the online retrieval functions and writes only inside the demo
    workspace.  Each selected calendar date produces exactly one summary row.
    """
    source = Path(project_dir or repository_root()).resolve()
    workspace = prepare_demo_workspace(source, workspace_dir, reset=reset)
    mode = str(mode).strip().upper()
    if mode not in {"REPLAY", "LIVE"}:
        raise ValueError("mode must be REPLAY or LIVE")

    start = _parse_date(start_date)
    end = _parse_date(end_date)
    if start > end:
        raise ValueError("START_DATE must not be later than END_DATE")

    calendar = pd.read_excel(
        workspace / "China_HK_Trading_Calendar.xlsx", sheet_name="Daily Calendar"
    )
    calendar["Date"] = pd.to_datetime(calendar["Date"]).dt.normalize()
    selected = calendar.loc[calendar["Date"].between(start, end)].copy()
    expected_days = (end - start).days + 1
    if len(selected) != expected_days:
        raise ValueError(
            "The selected range is not fully covered by the trading calendar"
        )

    ingestion.configure_project(workspace)
    replay = _load_replay_rows(source) if mode == "REPLAY" else pd.DataFrame()
    summaries = []
    saved_daily_rows = []

    for day in pd.date_range(start, end, freq="D"):
        day_text = day.strftime("%Y%m%d")
        calendar_row = selected.loc[selected["Date"] == day].iloc[0]
        hk_open = bool(calendar_row["HK_Open"])
        china_open = bool(calendar_row["China_Open"])

        try:
            if mode == "LIVE":
                daily = ingestion.process_daily_data(day_text, save_result=True).iloc[0]
            else:
                daily = _replay_daily_row(day, calendar_row, replay)
                if hk_open:
                    action, conflict = ingestion.update_temporary(daily, True)
                    if conflict:
                        raise RuntimeError("Replay row conflicts with Temporary Input")
                    daily["Temporary_Action"] = action
                saved_daily_rows.append(dict(daily))

            dtd = process_daily_dtd(
                day_text,
                confirmed_file=workspace / "vanke.xlsx",
                temporary_input_file=workspace / "vanke_dtd_temporary_data.xlsx",
                temporary_output_file=workspace / "temporary_output.xlsx",
                calendar_file=workspace / "China_HK_Trading_Calendar.xlsx",
                company=company,
            )
            summaries.append(
                {
                    "Date": day.strftime("%Y-%m-%d"),
                    "China_Open": china_open,
                    "HK_Open": hk_open,
                    "Market_Case": daily.get("Market_Case"),
                    "Input_QC": daily.get("Quality_Status"),
                    "Temporary_Input_Action": daily.get("Temporary_Action"),
                    "DTD_Status": dtd.status,
                    "DTD": dtd.dtd,
                    "Message": dtd.message,
                }
            )
        except Exception as error:
            summaries.append(
                {
                    "Date": day.strftime("%Y-%m-%d"),
                    "China_Open": china_open,
                    "HK_Open": hk_open,
                    "Market_Case": _market_case(china_open, hk_open),
                    "Input_QC": "BLOCKED",
                    "Temporary_Input_Action": "NOT_WRITTEN",
                    "DTD_Status": type(error).__name__,
                    "DTD": None,
                    "Message": str(error),
                }
            )

    if mode == "REPLAY":
        _write_replay_datalog(workspace, saved_daily_rows)

    temporary_input = pd.read_excel(
        workspace / "vanke_dtd_temporary_data.xlsx", sheet_name="Input"
    )
    temporary_dtd_path = workspace / "temporary_output.xlsx"
    temporary_dtd = (
        pd.read_excel(temporary_dtd_path, sheet_name="Output")
        if temporary_dtd_path.exists()
        else pd.DataFrame(columns=["Comp_no", "Date", "DTD"])
    )
    confirmed_input = pd.read_excel(workspace / "vanke.xlsx", sheet_name="Input")
    updated_clean_input = combine_input_history(
        confirmed_input, temporary_input, company=company
    )[INPUT_COLUMNS].copy()
    updated_clean_input["Date"] = (
        pd.to_datetime(updated_clean_input["Date"]).dt.strftime("%Y%m%d").astype(int)
    )
    return RangeRunResult(
        start_date=start.strftime("%Y-%m-%d"),
        end_date=end.strftime("%Y-%m-%d"),
        mode=mode,
        workspace=workspace,
        daily_results=pd.DataFrame(summaries),
        temporary_input=temporary_input,
        temporary_dtd=temporary_dtd,
        updated_clean_input=updated_clean_input,
    )


def _combine_results(daily: pd.Series, dtd: DailyProcessingResult) -> PipelineResult:
    return PipelineResult(
        date=pd.Timestamp(daily["Date"]).strftime("%Y-%m-%d"),
        market_case=str(daily["Market_Case"]),
        quality_status=str(daily["Quality_Status"]),
        temporary_action=str(daily["Temporary_Action"]),
        dtd_status=dtd.status,
        dtd=dtd.dtd,
        wrote_dtd_output=dtd.wrote_output,
        message=dtd.message,
    )


def _parse_date(value) -> pd.Timestamp:
    text = str(value).strip().removesuffix(".0")
    value = pd.to_datetime(
        text, format="%Y%m%d" if len(text) == 8 and text.isdigit() else None
    )
    return pd.Timestamp(value).normalize()


def _market_case(china_open: bool, hk_open: bool) -> str:
    if china_open and hk_open:
        return "BOTH_OPEN"
    if china_open:
        return "CHINA_ONLY"
    if hk_open:
        return "HK_ONLY"
    return "BOTH_CLOSED"


def _create_empty_temporary_input(workspace: Path) -> None:
    workbook = load_workbook(workspace / "vanke.xlsx")
    sheet = workbook["Input"]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for other in list(workbook.worksheets):
        if other.title != "Input":
            workbook.remove(other)
    workbook.save(workspace / "vanke_dtd_temporary_data.xlsx")


def _load_replay_rows(source: Path) -> pd.DataFrame:
    path = repository_data(source).replay_datalog
    if not path.exists():
        raise FileNotFoundError("Vanke_Daily_Datalog.xlsx is required for REPLAY mode")
    replay = pd.read_excel(path, sheet_name="Daily_Result")
    replay["Date"] = pd.to_datetime(replay["Date"]).dt.normalize()
    return replay.sort_values("Date").drop_duplicates("Date", keep="last")


def _replay_daily_row(day, calendar_row, replay: pd.DataFrame) -> dict:
    source_row = replay.loc[replay["Date"] == day]
    china_open = bool(calendar_row["China_Open"])
    hk_open = bool(calendar_row["HK_Open"])
    if not source_row.empty:
        return source_row.iloc[-1].to_dict()
    if hk_open:
        raise ValueError(
            f"No saved source snapshot exists for {day.date()}; use LIVE mode or choose a replay date"
        )
    return {
        "Date": day,
        "China_Open": china_open,
        "HK_Open": hk_open,
        "Market_Case": _market_case(china_open, hk_open),
        "Quality_Status": "PASS_EXPECTED_HK_CLOSED",
        "Ready_For_DTD": False,
        "Bug_Flag": False,
        "Review_Flag": False,
        "Temporary_Action": "NOT_ELIGIBLE_HK_CLOSED",
    }


def _write_replay_datalog(workspace: Path, rows) -> None:
    daily = pd.DataFrame(rows)
    with pd.ExcelWriter(
        workspace / "Vanke_Daily_Datalog.xlsx", engine="openpyxl"
    ) as writer:
        daily.to_excel(writer, sheet_name="Daily_Result", index=False)
        pd.DataFrame(
            columns=[
                "Run_ID",
                "Input_Date",
                "API",
                "Attempt",
                "Attempt_Status",
                "Attempt_Time",
                "Error_Type",
                "Error_Message",
            ]
        ).to_excel(writer, sheet_name="API_Attempt_Log", index=False)
